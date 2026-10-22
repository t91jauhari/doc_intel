"""Account Opening startup, pinned flow, SOP, comparison and XLSX regression coverage."""

import copy
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.account_opening_seed import RELEASE, seed_account_opening
from app.account_opening_sop import evaluate_rule, score, sources, CATEGORIES
from app.config_workbook import export_config, import_config
from app.dataset_comparison import ACCOUNT_DATASET, dataset_for_case
from app.db import Case, Config, Report, SessionLocal
from app.domain import default_config, latest_config
from app.transaction_flows import scope_configuration
from app.kertas_kerja import workbook
from test_platform import upload, publish_fixture, process_all, PDF

RULES = json.loads(
    (
        Path(__file__).resolve().parents[1]
        / "app/config_data/account_opening_rules.json"
    ).read_text()
)


def check(rule_id, documents, context=None):
    rule = next(r for r in RULES if r["id"] == rule_id)
    return evaluate_rule(rule, documents, default_config(), context or {})


def doc(kind, **values):
    return {
        "id": kind,
        "doc_type": kind,
        "status": "COMPLETED",
        "normalized": values,
        "validation": {"fields": {}},
    }


def test_seed_is_idempotent_and_forex_configuration_is_isolated(signed):
    with SessionLocal() as db:
        before = latest_config(db, "sqe")
        version, payload = before.version, copy.deepcopy(before.payload)
    assert RELEASE in payload["seed_releases"]
    assert len(payload["profiles"]) == 12
    assert len(payload["templates"]["account_opening"]["sop_rule_ids"]) == 206
    assert not any(
        r["id"].startswith("AO_")
        for r in scope_configuration(payload, "forex")["sop_rules"]
    )
    assert all(
        r["id"].startswith("AO_")
        for r in scope_configuration(payload, "account_opening")["sop_rules"]
    )
    seed_account_opening()
    with SessionLocal() as db:
        assert latest_config(db, "sqe").version == version


def test_existing_custom_account_flow_and_case_pin_survive_seed(signed):
    with SessionLocal.begin() as db:
        old = latest_config(db, "sqe")
        payload = copy.deepcopy(old.payload)
        payload["seed_releases"].remove(RELEASE)
        for name in (
            "join_visit",
            "edd",
            "account_opening_form",
            "signature_difference",
            "company_reference",
        ):
            payload["profiles"].pop(name)
            payload["profile_variants"].pop(name)
        payload["sop_rules"] = [
            r for r in payload["sop_rules"] if not r["id"].startswith("AO_")
        ]
        payload["templates"]["account_opening_reference"] = {
            "label": "User reference",
            "required": [],
            "sop_rule_ids": [],
        }

        custom = {
            "label": "Custom Account Opening",
            "required": ["ektp"],
            "sop_rule_ids": [],
            "enabled": True,
        }
        payload["templates"]["account_opening"] = custom
        db.add(
            Config(
                org="sqe", version=old.version + 1, payload=payload, published_by="test"
            )
        )
    response = upload(signed, names=("Synthetic_ektp.pdf",), template="account_opening")
    case_id = response.json()["case_ids"][0]
    with SessionLocal() as db:
        pinned = db.get(Case, case_id).config_id
    seed_account_opening()
    with SessionLocal() as db:
        current = latest_config(db, "sqe")
        assert current.payload["templates"]["account_opening"] == custom
        assert (
            current.payload["templates"]["account_opening_reference_2"]["dataset_id"]
            == ACCOUNT_DATASET["id"]
        )
        assert db.get(Case, case_id).config_id == pinned
        assert (
            current.payload["templates"]["account_opening_reference"]["label"]
            == "User reference"
        )
        assert len(current.payload["profiles"]) == 12


def test_reference_provenance_scope_and_workbook_roundtrip(signed):
    assert len(ACCOUNT_DATASET["rows"]) == 210 and len(ACCOUNT_DATASET["samples"]) == 9
    assert any("Sri" in note for note in ACCOUNT_DATASET["notes"])
    assert all(r["source_row"] and r["source_sheet"] for r in ACCOUNT_DATASET["rows"])
    assert "AO:SOP System:12" not in {s["source_ref"] for s in sources()}
    cfg = signed.get("/api/config").json()["payload"]
    imported, errors = import_config(export_config(cfg), cfg)
    assert not errors and imported["templates"] == cfg["templates"]
    by_id = {r["id"]: r for r in imported["sop_rules"]}
    for rule in cfg["sop_rules"]:
        if rule["id"].startswith("AO_"):
            assert all(
                by_id[rule["id"]][key] == ("" if value is None else value)
                for key, value in rule.items()
            )
    catalog = signed.get("/api/config/catalog").json()
    assert "OBSERVATIONS" in catalog["categories"]
    assert any(
        f["path"] == "account_opening_form.customer_signature"
        for f in catalog["fields"]
    )


@pytest.mark.parametrize(
    "value,result",
    [
        ("Risiko Tinggi", "FAIL"),
        ("Risiko Rendah", "NOT_APPLICABLE"),
        ("unclear", "REVIEW"),
    ],
)
def test_conditional_edd_requires_observed_trigger(value, result):
    assert (
        check(
            "AO_EDD_REQUIRED",
            [
                doc(
                    "account_opening_form",
                    customer_name="Synthetic",
                    risk_category=value,
                )
            ],
        )["result"]
        == result
    )


def test_external_and_unspecified_document_requirements_never_fabricate_pass():
    form = doc(
        "account_opening_form", customer_name="Synthetic", acting_for="Beneficial Owner"
    )
    assert check("AO_BO_REQUIRED", [form])["result"] == "REVIEW"
    assert check("AO_SURVEY", [form])["result"] == "REVIEW"
    assert check("AO_STAMP_CONDITION", [form])["result"] == "REVIEW"


def test_cross_customer_identifier_gender_and_nik_birth_checks():
    ktp = doc(
        "ektp",
        name="Synthetic Person",
        nik="1234564106930001",
        birth_date="01-06-1993",
        gender="Perempuan",
    )
    form = doc(
        "account_opening_form",
        customer_name="Different Person",
        gender="Wanita",
        identity_number="1234564106930001",
    )
    assert check("AO_SYSTEM_2", [ktp, form])["result"] == "FAIL"
    assert check("AO_SYSTEM_1", [ktp, form])["result"] == "PASS"
    assert check("AO_SYSTEM_5", [ktp, form])["result"] == "PASS"
    assert check("AO_NIK_BIRTH", [ktp])["result"] == "PASS"
    ktp["normalized"]["gender"] = "unclear"
    assert check("AO_NIK_BIRTH", [ktp])["result"] == "REVIEW"
    form["normalized"]["identity_number"] = "1234564106930002"
    assert check("AO_SYSTEM_1", [ktp, form])["result"] == "FAIL"


def test_account_scores_do_not_turn_unresolved_into_pass():
    findings = [{"category": c, "eligible": True, "result": "PASS"} for c in CATEGORIES]
    assert score(findings)["total"] == 1
    findings[-1]["result"] = "REVIEW"
    assert score(findings)["total_b"] is None and score(findings)["total"] is None
    findings[0]["result"] = "FAIL"
    assert score(findings)["total"] == 0


def test_fixture_pipeline_flow_dataset_and_account_report(signed):
    publish_fixture(signed)
    current = signed.get("/api/config").json()
    current["payload"]["maker_checker"] = False
    assert (
        signed.post(
            "/api/config/publish",
            json={
                "expected_version": current["version"],
                "payload": current["payload"],
            },
        ).status_code
        == 200
    )
    names = (
        "Synthetic_account opening.pdf",
        "Synthetic_ektp.pdf",
        "Synthetic_riplay.pdf",
        "Synthetic_edd.pdf",
    )
    branch = signed.get("/api/branches").json()[0]["id"]
    response = signed.post(
        "/api/uploads",
        data={"branch_id": branch, "template": "account_opening"},
        files=[
            ("files", (name, PDF + name.encode(), "application/pdf")) for name in names
        ],
        headers={"Idempotency-Key": "account-test"},
    )
    assert response.status_code == 202, response.text
    case_id = response.json()["case_ids"][0]
    process_all()
    detail = signed.get(f"/api/cases/{case_id}").json()
    assert {d["doc_type"] for d in detail["documents"]} == {
        "account_opening_form",
        "ektp",
        "riplay",
        "edd",
    }
    assert all(d["status"] == "COMPLETED" for d in detail["documents"])
    assert (
        signed.get(f"/api/dataset-sample?case_id={case_id}").json()["id"]
        == ACCOUNT_DATASET["id"]
    )
    assert signed.get("/api/dataset-sample?case_id=unknown").status_code == 404
    with SessionLocal() as db:
        case = db.get(Case, case_id)
        assert dataset_for_case(db, case)["id"] == ACCOUNT_DATASET["id"]
        snapshot = db.scalar(
            select(Report)
            .where(Report.case_id == case_id)
            .order_by(Report.version.desc())
        ).snapshot
    assert all(f["rule_id"].startswith("AO_") for f in snapshot["findings"])
    book = load_workbook(io.BytesIO(workbook([snapshot])))
    sheet = book["Account Opening Kertas Kerja"]
    assert (
        sheet["N11"].value == "Officer Review & Approval"
        and sheet["R11"].value == "Observations"
    )
    assert (
        sheet["E12"].value is None
    )  # Existing account is not the newly opened account.
    assert sheet["K12"].value and sheet.print_area.endswith(f"$U${sheet.max_row}")
    assert {
        "SOP Details",
        "Extracted Fields",
        "Visual Evidence",
        "Report Trace",
    } <= set(book.sheetnames)
    by_type = {d["doc_type"]: d["id"] for d in detail["documents"]}
    review = {
        "rule_id": "AO_CUSTOMER_SIGNATURE_COMPARE",
        "result": "PASS",
        "reason": "Compared both visible signing regions",
        "observations": [
            {
                "document_id": by_type["account_opening_form"],
                "page": 1,
                "region": "Customer signing block",
            }
        ],
    }
    assert (
        signed.post(f"/api/cases/{case_id}/sop-review", json=review).status_code == 422
    )
    review["observations"].append(
        {
            "document_id": by_type["riplay"],
            "page": 4,
            "region": "Customer signing block",
        }
    )
    assert (
        signed.post(f"/api/cases/{case_id}/sop-review", json=review).status_code == 202
    )
    process_all()
    rejected = signed.put(
        f"/api/cases/{case_id}/comparison",
        json={
            "sample_id": "1",
            "expected_revision": 0,
            "edits": {"forex-row-3": "wrong dataset"},
            "reason": "Test isolation",
        },
    )
    assert rejected.status_code == 422
    comparison = signed.put(
        f"/api/cases/{case_id}/comparison",
        json={
            "sample_id": "1",
            "expected_revision": 0,
            "reason": "Independent reference comparison",
        },
    )
    assert comparison.status_code == 200 and len(comparison.json()["rows"]) == 210
    assert (
        signed.put(
            f"/api/cases/{case_id}/context",
            json={
                "values": {"new_account_number": "00AB"},
                "reason": "Test annotation",
                "evidence": "Opening source page 1",
            },
        ).status_code
        == 422
    )
    assert (
        signed.put(
            f"/api/cases/{case_id}/context",
            json={
                "values": {"new_account_number": "00123456789"},
                "reason": "Confirmed new account from source",
                "evidence": "Opening confirmation page 1",
            },
        ).status_code
        == 200
    )
    process_all()
    with SessionLocal() as db:
        annotated = db.scalar(
            select(Report)
            .where(Report.case_id == case_id)
            .order_by(Report.version.desc())
        ).snapshot
    assert (
        load_workbook(io.BytesIO(workbook([annotated]))).worksheets[0]["E12"].value
        == "00123456789"
    )
    from PIL import Image
    from app.storage import store

    picture = io.BytesIO()
    Image.new("RGB", (24, 24), "red").save(picture, format="PNG")
    key = "sqe/test/account-export.png"
    store().put(key, picture.getvalue(), "image/png")
    snapshot["documents"][0]["validation"]["visual_evidence"] = [
        {
            "field": "customer_signature",
            "page_number": 1,
            "observed_present": True,
            "available": True,
            "review_status": "PENDING",
            "source_sha256": "test",
            "revision": "test",
            "storage_key": key,
        }
    ]
    forex = copy.deepcopy(snapshot)
    forex["template"] = "forex"
    forex["configuration"] = scope_configuration(forex["configuration"], "forex")
    mixed = load_workbook(io.BytesIO(workbook([forex, snapshot])))
    assert "Kertas Kerja" in mixed and "AO Account Opening Kertas Kerja" in mixed
    assert len(mixed["AO Visual Evidence"]._images) == 1
    assert (
        mixed["AO Account Opening Kertas Kerja"]["A11"].fill.fgColor.rgb == "000F766E"
    )


@pytest.mark.parametrize(
    "field",
    [
        "visitor_signature",
        "edd_officer_signature",
        "customer_signature",
        "signature_one",
        "signatory_signature",
        "confirmation_stamp",
        "meterai",
    ],
)
def test_new_visual_fields_use_actual_pdf_regions(field):
    from types import SimpleNamespace
    from app.evidence import build_evidence
    from app.storage import original_key, store
    import fitz

    pdf = fitz.open()
    page = pdf.new_page(width=100, height=100)
    page.draw_rect(fitz.Rect(20, 20, 40, 40), fill=(1, 0, 0))
    source = pdf.tobytes()
    pdf.close()
    document = SimpleNamespace(
        org="sqe",
        branch_id="test",
        id="account-" + field,
        sha256="test",
        provider_state={},
    )
    store().put(original_key(document), source, "application/pdf")
    evidence = build_evidence(
        document,
        {
            "id": "synthetic",
            "extract_result": {field: {"present": True}},
            "extract_metadata": {
                "visual_evidence": {
                    field: {
                        "page_number": 1,
                        "bbox": [20, 20, 40, 40],
                        "units": "points",
                    }
                }
            },
        },
    )
    assert (
        len(evidence) == 1
        and evidence[0]["available"]
        and evidence[0]["field"] == field
    )
    assert store().read(evidence[0]["storage_key"]).startswith(b"\x89PNG")
