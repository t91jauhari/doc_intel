import io
from copy import deepcopy

from app.case_processing import process_case_run
from app.config_workbook import export_config, import_config
from app.db import CaseRun, Job, SessionLocal
from app.domain import default_config, validate_config
from app.sop import CATEGORIES, catalog, evaluate, fields, score
from openpyxl import load_workbook
from sqlalchemy import select
from test_platform import PDF, process_all, publish_fixture, upload


def doc(kind, values, **extra):
    return {
        "id": kind,
        "doc_type": kind,
        "status": "COMPLETED",
        "normalized": values,
        "validation": {},
        "overrides": {},
        **extra,
    }


def rule(cfg, id="SYS_01_NAME"):
    return deepcopy(next(r for r in cfg["sop_rules"] if r["id"] == id))


def test_scope_is_only_yes_and_all_sop3():
    sources = catalog("scope")
    assert len(sources) == 29
    assert {s["source_ref"] for s in sources if s["scope_basis"] == "SYSTEM_SOP3"} == {
        f"SOP3:{i}" for i in range(1, 16)
    }
    assert {
        s["source_ref"] for s in sources if s["source_ref"].startswith("SOP1:")
    } == {f"SOP1:{i} | SOP2:{i}" for i in [2, 4, 5, 6, 7, 9]}
    assert set(r["source_ref"] for r in default_config()["sop_rules"]) == {
        s["source_ref"] for s in sources
    }
    assert "application.identity_number" not in {f["path"] for f in fields()}


def test_invalid_external_source_operator_and_workflow_rejected():
    for change in [
        ("source_ref", "SOP2:1"),
        ("operator", "cctv"),
        ("path", "application.identity_number"),
    ]:
        cfg = default_config()
        cfg["sop_rules"][0][change[0]] = change[1]
        try:
            validate_config(cfg)
        except ValueError:
            pass
        else:
            assert False, change
    cfg = default_config()
    cfg["workflow"][2]["order"] = 8
    try:
        validate_config(cfg)
    except ValueError:
        pass
    else:
        assert False


def test_import_roundtrip_and_atomic_row_errors():
    cfg = default_config()
    data = export_config(cfg)
    draft, errors = import_config(data, cfg)
    assert not errors and draft is not None
    assert draft["profiles"] == cfg["profiles"] and draft["workflow"] == cfg["workflow"]
    w = load_workbook(io.BytesIO(data))
    w["SOP Rules"]["G2"] = "nonexistent"
    out = io.BytesIO()
    w.save(out)
    draft, errors = import_config(out.getvalue(), cfg)
    assert draft is None and any(
        e["sheet"] == "SOP Rules" and e["row"] == 2 for e in errors
    )
    assert cfg == default_config()
    w = load_workbook(io.BytesIO(data))
    w["Workflow"]["F2"] = '=HYPERLINK("http://evil")'
    out = io.BytesIO()
    w.save(out)
    assert import_config(out.getvalue(), cfg)[0] is None


def test_import_does_not_publish(signed):
    version = signed.get("/api/config").json()["version"]
    data = signed.get("/api/config/export")
    r = signed.post(
        "/api/config/import",
        files={
            "file": (
                "config.xlsx",
                data.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200 and r.json()["valid"]
    assert signed.get("/api/config").json()["version"] == version
    assert (
        signed.post(
            "/api/config/import",
            files={"file": ("bad.xlsx", b"invalid", "application/octet-stream")},
        ).json()["valid"]
        is False
    )


def test_name_currency_dates_and_low_confidence():
    cfg = default_config()
    cfg["sop_rules"] = [rule(cfg)]
    docs = [
        doc("ektp", {"name": "Rina Putri"}),
        doc("application", {"remitter": {"name": " RINA  PUTRI "}}),
    ]
    assert evaluate(docs, cfg, {})[0]["result"] == "PASS"
    docs[1]["normalized"]["remitter"]["name"] = "Other Name"
    assert evaluate(docs, cfg, {})[0]["result"] == "FAIL"
    docs[1]["normalized"]["remitter"]["name"] = "Rina Putri"
    docs[0]["validation"] = {
        "fields": {"name": {"issues": [{"code": "LOW_EXTRACTION_CONFIDENCE"}]}}
    }
    assert evaluate(docs, cfg, {})[0]["result"] == "FAIL"
    cfg["sop_rules"] = [rule(default_config(), "SYS_08_AMOUNT")]
    docs = [
        doc("application", {"amount": "1000.00", "currency": "USD"}),
        doc("lld", {"transaction_value": {"amount": "1000", "currency": "EUR"}}),
    ]
    assert evaluate(docs, cfg, {})[0]["result"] == "FAIL"
    docs[1]["normalized"]["transaction_value"]["currency"] = "USD"
    assert evaluate(docs, cfg, {})[0]["result"] == "PASS"
    cfg["sop_rules"] = [rule(default_config(), "SYS_13_DATE")]
    docs = [
        doc("application", {"transaction_date": "06-10-2026"}),
        doc("lld", {"transaction_date": "2026-10-06"}),
    ]
    assert evaluate(docs, cfg, {})[0]["result"] == "PASS"  # Same calendar date.


def test_valas2_applicability_and_explicit_invoice_basis():
    cfg = default_config()
    cfg["sop_rules"] = [rule(cfg, "SYS_10_AMOUNT"), rule(cfg, "SYS_12_PURPOSE")]
    docs = [
        doc("valas_declaration_2", {"currency": "USD"}),
        doc("lld", {"transaction_value": {"amount": "60", "currency": "USD"}}),
    ]
    assert {f["result"] for f in evaluate(docs, cfg, {})} == {"NOT_APPLICABLE"}
    cfg["sop_rules"] = [rule(default_config(), "SYS_11_INVOICE")]
    docs += [
        doc(
            "invoice",
            {
                "final_total_amount": {"amount": "100", "currency": "USD"},
                "claim_note": {"amount": "60", "currency": "USD"},
            },
        )
    ]
    assert evaluate(docs, cfg, {})[0]["result"] == "FAIL"
    assert (
        evaluate(docs, cfg, {"payment_basis": "PROGRESS_CLAIM"})[0]["result"] == "PASS"
    )
    assert evaluate(docs, cfg, {"payment_basis": "FULL_INVOICE"})[0]["result"] == "FAIL"


def test_unknown_context_fails_not_false_na():
    cfg = default_config()
    cfg["sop_rules"] = [rule(cfg, "SYS_02_NIK")]
    docs = [
        doc("ektp", {"nik": "1234567890123456"}),
        doc("valas_declaration_1", {"identity_number": "1234567890123456"}),
    ]
    assert evaluate(docs, cfg, {})[0]["result"] == "FAIL"
    assert evaluate(docs, cfg, {"identity_type": "NIK"})[0]["result"] == "PASS"
    assert (
        evaluate(docs, cfg, {"identity_type": "PASSPORT"})[0]["result"]
        == "NOT_APPLICABLE"
    )


def test_visual_presence_does_not_prove_signature_consistency():
    cfg = default_config()
    cfg["sop_rules"] = [rule(cfg, "SYS_05_SIG")]
    sig = {
        "present": True,
        "non_text": True,
        "complete": True,
        "page_number": 4,
        "evidence": "visible strokes",
    }
    assert (
        evaluate(
            [
                doc("riplay", {"customer_signature": sig}),
                doc("application", {"customer_signature": sig}),
            ],
            cfg,
            {},
        )[0]["result"]
        == "REVIEW"
    )
    cfg["sop_rules"] = [rule(default_config(), "YES_RIPLAY_PAGES")]
    observations = {
        "page_count": 4,
        "pages": [
            {
                "page_number": i,
                "footer_present": True,
                "non_text": True,
                "complete": True,
                "evidence": "bottom-right",
            }
            for i in range(1, 5)
        ],
    }
    assert (
        evaluate([doc("riplay", {"page_initials": observations})], cfg, {})[0]["result"]
        == "PASS"
    )
    observations["pages"][1]["footer_present"] = False
    assert (
        evaluate([doc("riplay", {"page_initials": observations})], cfg, {})[0]["result"]
        == "FAIL"
    )


def test_score_excludes_qa_and_zero_dominates_unknown():
    findings = [
        {"eligible": True, "category": CATEGORIES[0], "result": "FAIL"},
        {"eligible": True, "category": CATEGORIES[1], "result": "REVIEW"},
        {"eligible": False, "category": CATEGORIES[2], "result": "FAIL"},
    ]
    scores = score(findings)
    assert (
        scores["total_a"] == 0
        and scores["total"] == 0
        and scores[CATEGORIES[2]] is None
    )
    assert score([])["total"] is None


def test_separate_case_jobs_frozen_snapshot_and_no_paid_rerun(signed, monkeypatch):
    publish_fixture(signed)
    case_id = upload(signed, names=("rina_ektp.pdf",), template="custom").json()[
        "case_ids"
    ][0]
    from app.processor import process_job

    with SessionLocal() as db:
        job_id = db.scalar(select(Job.id))
    process_job(job_id)
    detail = signed.get("/api/cases/" + case_id).json()
    assert detail["status"] == "SOP_QUEUED" and not detail["reports"]
    old_run = detail["sop_runs"][0]["id"]
    ctx = signed.put(
        "/api/cases/" + case_id + "/context",
        json={
            "values": {"customer_type": "INDIVIDUAL"},
            "reason": "Evidence checked",
            "evidence": "eKTP first page",
        },
    )
    assert ctx.status_code == 200
    process_case_run(old_run)
    with SessionLocal() as db:
        assert db.get(CaseRun, old_run).status == "SUPERSEDED"
    process_all()
    detail = signed.get("/api/cases/" + case_id).json()
    before = detail["reports"][0]["id"]
    monkeypatch.setattr(
        "app.processor.run_provider",
        lambda *args: (_ for _ in ()).throw(AssertionError("paid job")),
    )
    assert signed.post("/api/cases/" + case_id + "/sop-rerun").status_code == 202
    process_all()
    assert signed.get("/api/cases/" + case_id).json()["reports"][0]["id"] == before


def test_confirmed_fail_finality_separate_from_compliance(signed, monkeypatch):
    # Isolate release scoring; visual evidence gating is tested separately.
    monkeypatch.setattr("app.processor.build_evidence", lambda *_: [])
    cfg = signed.get("/api/config").json()
    cfg["payload"]["provider"] = "fixture"
    r = rule(cfg["payload"], "YES_APP_REQUIRED")
    r["path"] = "application.notes"
    cfg["payload"]["sop_rules"] = [r]
    for flow in cfg["payload"]["templates"].values():
        flow["sop_rule_ids"] = None
    assert (
        signed.post(
            "/api/config/publish",
            json={"expected_version": cfg["version"], "payload": cfg["payload"]},
        ).status_code
        == 200
    )
    case_id = upload(signed, names=("poko_transaction.pdf",), template="custom").json()[
        "case_ids"
    ][0]
    process_all()
    # Set an explicit blank entered field, then revalidate stored raw.
    detail = signed.get("/api/cases/" + case_id).json()
    docid = detail["documents"][0]["id"]
    signed.post(
        "/api/documents/" + docid + "/corrections",
        json={
            "field": "notes",
            "value": "",
            "reason": "Original notes region inspected blank",
        },
    )
    process_all()
    detail = signed.get("/api/cases/" + case_id).json()
    assert detail["validation_status"] == "FAIL" and detail["status"] == "READY"
    report = detail["reports"][0]
    assert signed.post("/api/reports/" + report["id"] + "/approve").status_code == 200


def test_branch_report_layout_scores_and_identifier_text(signed):
    publish_fixture(signed)
    upload(
        signed,
        names=("poko_transaction.pdf", "poko_lld.pdf", "poko_invoice.pdf"),
        key="layout1",
    )
    process_all()
    upload(signed, names=("rina_transaction.pdf",), key="layout2", template="custom")
    process_all()
    branch = signed.get("/api/branches").json()[0]["id"]
    r = signed.get("/api/reports/kertas-kerja?branch_id=" + branch)
    assert r.status_code == 200
    w = load_workbook(io.BytesIO(r.content), data_only=False)
    s = w["Kertas Kerja"]
    assert s["C12"].value == s["C13"].value == 1
    assert isinstance(s["E12"].value, str) and s["E12"].value.startswith("0")
    assert s.column_dimensions["S"].hidden
    assert all(
        s.cell(row, col).data_type != "f"
        for row in range(12, 20)
        for col in range(1, 21)
    )
    assert "SOP Details" in w and "Document Quality" in w
    assert signed.get(
        "/api/reports/kertas-kerja?branch_id=another-branch"
    ).status_code in {403, 404}


def test_review_requires_related_documents_and_both_signature_regions(signed):
    cfg = signed.get("/api/config").json()
    cfg["payload"]["provider"] = "fixture"
    cfg["payload"]["sop_rules"] = [rule(cfg["payload"], "SYS_05_SIG")]
    for flow in cfg["payload"]["templates"].values():
        flow["sop_rule_ids"] = None
    assert (
        signed.post(
            "/api/config/publish",
            json={"expected_version": cfg["version"], "payload": cfg["payload"]},
        ).status_code
        == 200
    )
    branch = signed.get("/api/branches").json()[0]["id"]
    case_id = signed.post(
        "/api/uploads",
        data={"branch_id": branch, "template": "custom"},
        files=[
            ("files", (name, PDF + name.encode(), "application/pdf"))
            for name in ["rina_riplay.pdf", "rina_transaction.pdf"]
        ],
        headers={"Idempotency-Key": "signature-review"},
    ).json()["case_ids"][0]
    process_all()
    detail = signed.get("/api/cases/" + case_id).json()
    docs = detail["documents"]
    ids = [d["id"] for d in docs]
    body = {
        "rule_id": "SYS_05_SIG",
        "result": "PASS",
        "reason": "Inspected supplied signing regions and observed consistent strokes",
        "observations": [
            {"document_id": ids[0], "page": 1, "region": "customer signing block"}
        ],
    }
    assert (
        signed.post("/api/cases/" + case_id + "/sop-review", json=body).status_code
        == 422
    )
    body["observations"].append(
        {"document_id": ids[1], "page": 4, "region": "customer signing block"}
    )
    body["report_id"] = "outdated-report"
    stale = signed.post("/api/cases/" + case_id + "/sop-review", json=body)
    assert stale.status_code == 409 and "report changed" in stale.json()["detail"]
    body["report_id"] = detail["reports"][0]["id"]
    assert detail["reports"][0]["snapshot"]["sop_summary"]["pending"] == 1

    assert (
        signed.post("/api/cases/" + case_id + "/sop-review", json=body).status_code
        == 202
    )
    process_all()
    detail = signed.get("/api/cases/" + case_id).json()
    f = detail["reports"][0]["snapshot"]["findings"][0]
    assert f["result"] == "PASS" and f["automated_result"] == "REVIEW"
    assert f["evidence"]["review"]["document_basis"]
    # Changed document values invalidate prior signature observations.
    signed.post(
        "/api/documents/" + ids[0] + "/corrections",
        json={
            "field": "customer_signature",
            "value": {
                "present": True,
                "non_text": True,
                "complete": False,
                "corrupted": False,
                "page_number": 4,
                "evidence": "Cropped signing area",
            },
            "reason": "Rechecked cropped document region",
        },
    )
    process_all()
    detail = signed.get("/api/cases/" + case_id).json()
    assert detail["reports"][0]["snapshot"]["findings"][0]["result"] == "REVIEW"
    # Context annotations explicitly discard old observation approvals.
    signed.put(
        "/api/cases/" + case_id + "/context",
        json={
            "values": {"customer_type": "INDIVIDUAL"},
            "reason": "Updated supplied context",
            "evidence": "Application page 1",
        },
    )
    assert not signed.get("/api/cases/" + case_id).json()["case_context"][
        "rule_reviews"
    ]
