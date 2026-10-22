import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.db import Case, Document, SessionLocal, User
from app.effective_corrections import extraction_value
from test_platform import process_all, publish_fixture, upload


def test_effective_correction_revalidates_without_paid_call(signed, monkeypatch):
    publish_fixture(signed)
    case_id = upload(signed, names=("Rina Putri_ektp.pdf",)).json()["case_ids"][0]
    process_all()
    detail = signed.get(f"/api/cases/{case_id}").json()
    doc = detail["documents"][0]
    assert doc["status"] == "COMPLETED"
    raw = doc["raw"]
    before = doc["normalized"]["name"]
    url = f'/api/documents/{doc["id"]}/effective-corrections'
    body = {"field": "name", "value": "Corrected Customer", "expected_value": before, "reason": "Corrected against source document"}
    response = signed.post(url, json=body)
    assert response.status_code == 202, response.text
    assert response.json()["normalized_preview"] == "Corrected Customer"
    assert signed.post(url, json=body).status_code == 409
    monkeypatch.setattr("app.processor.run_provider", lambda *a, **k: (_ for _ in ()).throw(AssertionError("No provider call for correction")))
    process_all()
    corrected = signed.get(f"/api/cases/{case_id}").json()["documents"][0]
    assert corrected["normalized"]["name"] == "Corrected Customer"
    assert corrected["raw"] == raw
    assert corrected["overrides"]["name"]["raw"] == "Corrected Customer"
    assert signed.post(url, json=body).status_code == 409
    assert signed.post(url, json={**body, "expected_value": "Corrected Customer", "value": "Corrected Customer"}).status_code == 422
    assert signed.post(url, json={**body, "expected_value": "Corrected Customer", "value": {"unexpected": 1}}).status_code == 422
    with SessionLocal.begin() as db:
        identity = signed.get("/api/auth/session").json()["user"]
        db.get(User, identity["id"]).role = "officer"
    assert signed.post(url, json=body).status_code == 403


def test_numeric_and_presence_adapters_preserve_observations():
    raw = {"raw": "100,00", "currency": "USD"}
    assert extraction_value("amount", "1234.50", raw, {"number_locale": "id_ID"}) == {"raw": "1234,50", "currency": "USD"}
    assert raw == {"raw": "100,00", "currency": "USD"}
    assert extraction_value("bank_approval", False, {"present": True, "stamp_present": True}, {}) == {"present": True, "stamp_present": False}


def test_one_excel_includes_all_accessible_comparisons_and_literal_values(signed):
    assert signed.get("/api/comparisons/export").status_code == 404
    case_ids = []
    for index, customer in enumerate(("First Customer", "Second Customer"), 1):
        case_id = upload(signed, names=(f"{customer}_ektp.pdf",), key=f"excel-{index}").json()["case_ids"][0]
        case_ids.append(case_id)
        with SessionLocal.begin() as db:
            doc = db.scalar(select(Document).where(Document.case_id == case_id))
            doc.doc_type = "ektp"
            doc.status = "COMPLETED" if index == 1 else "QUEUED"
            doc.normalized = {"name": customer, "nik": "0001234567890123"}
        saved = signed.put(f"/api/cases/{case_id}/comparison", json={"sample_id": str(index), "expected_revision": 0, "reason": "Selected sample reference", "edits": {"row-4": "=1+1", "row-79": {"name": "Example"}}})
        assert saved.status_code == 200
    response = signed.get("/api/comparisons/export")
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Summary", "Document results", "Field comparisons", "Reference edits", "Export information"]
    assert book["Summary"].max_row == 3
    assert book["Document results"].max_row == 17
    assert book["Document results"].cell(2, 8).value == 0
    assert book["Document results"].cell(2, 10).value == 10
    assert book["Field comparisons"].max_row == 327
    assert book["Summary"].freeze_panes == "A2"
    assert book["Field comparisons"].auto_filter.ref
    assert book["Summary"].cell(1, 1).fill.fgColor.rgb.endswith("167B70")
    assert any(row[10].value == "'=1+1" for row in list(book["Field comparisons"])[1:])
    assert any(row[12].value == "PENDING" for row in list(book["Field comparisons"])[1:])
    assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
    with SessionLocal.begin() as db:
        db.get(Case, case_ids[1]).org = "other-organization"
    restricted = load_workbook(io.BytesIO(signed.get("/api/comparisons/export").content))
    assert restricted["Summary"].max_row == 2
    assert all(row[1].value == case_ids[0] for row in list(restricted["Field comparisons"])[1:])
    signed.cookies.clear()
    assert signed.get("/api/comparisons/export").status_code == 401
