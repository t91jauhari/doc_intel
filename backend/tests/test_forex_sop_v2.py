"""Executable acceptance cases for the 15 source Forex checkpoints."""
from copy import deepcopy
import io
import pytest
from openpyxl import load_workbook
from app.domain import default_config, validate_config
from app.forex_sop import CheckError, address, money, package_summary, policy_for
from app.sop import evaluate
from app.config_workbook import export_config, import_config


def doc(kind, values, id=None):
    return {"id":id or kind,"doc_type":kind,"status":"COMPLETED","normalized":values,"validation":{},"overrides":{}}

@pytest.fixture
def sample():
    cfg=default_config()
    cfg["sop_rules"]=[r for r in cfg["sop_rules"] if r["source_ref"].startswith("SOP3:")]
    signature={"present":True,"complete":True,"non_text":True}
    docs=[
      doc("ektp",{"name":"Test Customer","address":"Jl. Melati No. 1","nik":"0012345678901234"}),
      doc("application",{"remitter":{"name":"Test Customer"},"beneficiary":{"name":"Test Supplier","address":"Jl Mawar 1"},"amount":"1000.00","currency":"USD","transaction_date":"10-09-2025","transaction_purpose":"Import goods","customer_signature":signature}),
      doc("valas_declaration_1",{"customer_name":"Test Customer","customer_address":"Jalan Melati 1","identity_number":"0012345678901234","usage_purpose":"Transaksi barang impor","usage_amount":{"amount":"1000.00","currency":"USD"}}),
      doc("lld",{"payer_name":"Test Customer","beneficiary_name":"Test Supplier","beneficiary_address":"Jl. Mawar No. 1","transaction_value":{"amount":"1000.00","currency":"USD"},"purpose":"Import goods","transaction_date":"2025-09-10"}),
      doc("invoice",{"invoice_number":"INV-001","to_name":"Test Customer","from_name":"Test Supplier","final_total_amount":{"amount":"1000.00","currency":"USD"},"claim_note":{"amount":"250.00","currency":"USD"}}),
      doc("riplay",{"customer_name":"Test Customer","document_date":"10-09-2025","customer_signature":signature}),
    ]
    return cfg,docs,{"identity_type":"NIK","payment_basis":"FULL_INVOICE"}


def results(sample):
    cfg,docs,context=sample
    return {f["rule_id"]:f for f in evaluate(docs,cfg,context)}


def test_complete_package_and_visual_denominator(sample):
    findings=list(results(sample).values())
    assert len(findings)==18
    assert sum(f["result"]=="PASS" for f in findings)==17
    assert next(f for f in findings if f["rule_id"]=="SYS_05_SIG")["result"]=="REVIEW"
    package=package_summary(findings)["forex"]
    assert package["coverage"]["complete"]
    assert (package["summary"]["earned"],package["summary"]["possible"],package["summary"]["pending"])==(14,15,1)
    for f in findings:
        if f["rule_id"]=="SYS_05_SIG": f["result"]="PASS"
    assert package_summary(findings)["forex"]["summary"]["earned"]==15
    findings[-1]["result"]="FAIL" # one child of checkpoint 7 fails the parent
    assert package_summary(findings)["forex"]["summary"]["earned"]==14


def test_exact_full_names_and_structured_readability(sample):
    cfg,docs,ctx=sample
    docs[0]["normalized"]["name"]={"raw":"TEST   CUSTOMER","complete":True}
    assert results(sample)["SYS_01_NAME"]["result"]=="PASS"
    docs[0]["normalized"]["name"]["complete"]=False
    assert results(sample)["SYS_01_NAME"]["reason_code"]=="UNREADABLE_SOURCE"
    docs[0]["normalized"]["name"]="Test Custome"
    assert results(sample)["SYS_01_NAME"]["result"]=="FAIL"
    docs[0]["normalized"]["nik"]=12345678901234
    assert results(sample)["SYS_02_NIK"]["reason_code"]=="INVALID_IDENTIFIER"


def test_address_aliases_preserve_units(sample):
    assert results(sample)["SYS_02_ADDRESS"]["result"]=="PASS"
    sample[1][0]["normalized"]["address"]="Jl Melati 2"
    assert results(sample)["SYS_02_ADDRESS"]["result"]=="FAIL"
    sample[1][0]["normalized"]["address"]="Jl Melati 1A"
    assert results(sample)["SYS_02_ADDRESS"]["result"]=="FAIL"

@pytest.mark.parametrize("value,expected",[("USD 100.817,90","100817.90"),("100,817.90","100817.90"),("100 817,90","100817.90"),("100817.90","100817.90")])
def test_money_formats(value,expected):
    assert str(money(value,"USD"))==expected

@pytest.mark.parametrize("value",["1,00,000.00","0","-2","NaN","1.234"])
def test_money_invalid_or_ambiguous(value):
    with pytest.raises(CheckError): money(value,"USD")


def test_currency_and_exact_amount(sample):
    sample[1][1]["normalized"]["amount"]="1.000,00"
    assert results(sample)["SYS_08_AMOUNT"]["result"]=="PASS"
    sample[1][1]["normalized"]["amount"]="1000.01"
    assert results(sample)["SYS_08_AMOUNT"]["result"]=="FAIL"
    sample[1][1]["normalized"].update(amount="1000",currency="EUR")
    assert results(sample)["SYS_08_AMOUNT"]["result"]=="FAIL"


def test_purpose_approved_codes_and_dates(sample):
    cfg,docs,ctx=sample
    docs[3]["normalized"]["purpose"]="012"
    assert results(sample)["SYS_12_PURPOSE"]["reason_code"]=="UNMAPPED_PURPOSE"
    cfg["forex_sop_policy"]["purpose_codes"]={"012":"Import goods"} # explicit test policy, not a shipped regulatory assertion
    assert results(sample)["SYS_12_PURPOSE"]["result"]=="PASS"
    docs[3]["normalized"]["transaction_date"]="09-10-2025"
    assert results(sample)["SYS_13_DATE"]["result"]=="FAIL"
    docs[3]["normalized"]["transaction_date"]="31-02-2025"
    assert results(sample)["SYS_13_DATE"]["reason_code"]=="INVALID_DATE"


def test_valas2_has_explicit_non_applicable_fields(sample):
    cfg,docs,ctx=sample
    docs[2]=doc("valas_declaration_2",{"customer_name":"Test Customer"})
    found=results(sample)
    assert found["SYS_06_PARTY_POLICY"]["result"]=="PASS"
    for key in ["SYS_02_NIK","SYS_02_NAME","SYS_02_ADDRESS","SYS_09_PURPOSE","SYS_10_AMOUNT","SYS_12_PURPOSE"]:
        assert found[key]["result"]=="NOT_APPLICABLE"
    ctx["threshold_state"]="ABOVE"
    assert results(sample)["SYS_10_AMOUNT"]["result"]=="FAIL"


def test_multiple_documents_require_binding(sample):
    cfg,docs,ctx=sample
    docs.append(doc("invoice",{"invoice_number":"INV-002"},"second"))
    assert results(sample)["SYS_11_INVOICE"]["reason_code"]=="AMBIGUOUS_DOCUMENT"
    ctx["invoice_document_id"]="invoice"
    assert results(sample)["SYS_11_INVOICE"]["result"]=="PASS"
    ctx["invoice_document_id"]="wrong"
    assert results(sample)["SYS_11_INVOICE"]["result"]=="FAIL"
    docs.append(doc("valas_declaration_2",{"customer_name":"Test Customer"}))
    assert results(sample)["SYS_06_PARTY_POLICY"]["reason_code"]=="AMBIGUOUS_DOCUMENT"
    ctx["declaration_document_id"]="valas_declaration_2"
    assert results(sample)["SYS_06_PARTY_POLICY"]["result"]=="PASS"


def test_invoice_identity_and_party_roles(sample):
    cfg,docs,ctx=sample
    assert results(sample)["SYS_14_LINK_POLICY"]["result"]=="PASS"
    docs[4]["normalized"]["to_name"]="Other customer"
    assert results(sample)["SYS_14_LINK_POLICY"]["result"]=="FAIL"
    docs[4]["normalized"]["to_name"]="Test Customer"
    docs[4]["normalized"]["invoice_number"]=None
    assert results(sample)["SYS_14_LINK_POLICY"]["reason_code"]=="MISSING_SOURCE_VALUE"


def test_supporting_invoice_and_quality(sample):
    cfg,docs,ctx=sample
    docs.append(deepcopy(docs[4]));docs[-1]["id"]="underlying"
    ctx.update(invoice_document_id="invoice",underlying_document_id="underlying")
    assert results(sample)["SYS_14_LINK_POLICY"]["result"]=="PASS"
    docs[-1]["normalized"]["invoice_number"]="INV-002"
    assert results(sample)["SYS_14_LINK_POLICY"]["result"]=="FAIL"
    ctx["underlying_document_id"]=""
    docs[4]["validation"]={"fields":{"to_name":{"issues":[{"code":"UNREADABLE_TEXT"}]}}}
    assert results(sample)["SYS_14_LINK_POLICY"]["reason_code"]=="UNVERIFIED_SOURCE"


def test_progress_partial_and_unknown_context(sample):
    cfg,docs,ctx=sample
    ctx["payment_basis"]="UNKNOWN"
    assert results(sample)["SYS_11_INVOICE"]["reason_code"]=="MISSING_PAYMENT_BASIS"
    ctx["payment_basis"]="PROGRESS_CLAIM";docs[3]["normalized"]["transaction_value"]["amount"]="250"
    assert results(sample)["SYS_11_INVOICE"]["result"]=="PASS"
    ctx["payment_basis"]="PARTIAL_PAYMENT"
    assert results(sample)["SYS_11_INVOICE"]["reason_code"]=="PAYMENT_EVIDENCE_REQUIRED"
    ctx.update(partial_payment_amount="250",partial_payment_currency="USD",annotation={"evidence":"Invoice page 1 payment schedule"})
    assert results(sample)["SYS_11_INVOICE"]["result"]=="PASS"
    ctx["partial_payment_amount"]="1500"
    assert results(sample)["SYS_11_INVOICE"]["reason_code"]=="INVALID_PAYMENT_BASIS"
    ctx.pop("identity_type")
    assert results(sample)["SYS_02_NIK"]["reason_code"]=="MISSING_CASE_CONTEXT"


def test_roundtrip_policy_and_legacy_workbook():
    cfg=default_config();cfg["forex_sop_policy"]["purpose_codes"]={"TEST":"Import goods"}
    draft,errors=import_config(export_config(cfg),cfg)
    assert not errors and draft["forex_sop_policy"]==cfg["forex_sop_policy"]
    book=load_workbook(io.BytesIO(export_config(cfg)))
    sheet=book["General Settings"]
    for row in sheet:
        if row[0].value=="forex_sop_policy":sheet.delete_rows(row[0].row);break
    output=io.BytesIO();book.save(output)
    assert not import_config(output.getvalue(),cfg)[1]
    cfg["forex_sop_policy"]["purpose_equivalences"]["other"]=["import goods"]
    with pytest.raises(ValueError):validate_config(cfg)


def test_missing_checkpoint_cannot_claim_full_package(sample):
    findings=[f for f in results(sample).values() if f["source_ref"]!="SOP3:14"]
    package=package_summary(findings)["forex"]
    assert not package["coverage"]["complete"]
    assert package["summary"]["failed"]==1


def test_runtime_adapter_preserves_extraction_configuration(sample):
    cfg,docs,ctx=sample
    for rule in cfg["sop_rules"]:
        if rule["id"] in {"SYS_06_PARTY_POLICY","SYS_14_LINK_POLICY"}:rule["operator"]="policy_review"
    before=deepcopy(cfg)
    found=results(sample)
    assert found["SYS_06_PARTY_POLICY"]["result"]=="PASS"
    assert found["SYS_14_LINK_POLICY"]["result"]=="PASS"
    assert cfg==before


def test_missing_compound_child_fails_package(sample):
    findings=[f for f in results(sample).values() if f["rule_id"]!="SYS_02_ADDRESS"]
    package=package_summary(findings)["forex"]
    assert not package["coverage"]["complete"]
    assert next(c for c in package["checkpoints"] if c["rule_id"]=="SOP3:2")["result"]=="FAIL"


def test_numeric_amount_has_no_localized_separator():
    assert str(money(1.234,"USD","id_ID"))=="1.234"
    with pytest.raises(CheckError): money("10 08 17,90","USD")


def test_context_document_selection_and_report_exports(signed):
    from test_platform import publish_fixture, upload, process_all
    publish_fixture(signed)
    case_id=upload(signed).json()["case_ids"][0]
    process_all()
    detail=signed.get("/api/cases/"+case_id).json()
    invoice=next(d for d in detail["documents"] if d["doc_type"]=="invoice")
    application=next(d for d in detail["documents"] if d["doc_type"]=="application")
    payload={"values":{"invoice_document_id":application["id"]},"reason":"Source invoice selected","evidence":"Uploaded invoice first page"}
    assert signed.put("/api/cases/"+case_id+"/context",json=payload).status_code==422
    payload["values"]={"invoice_document_id":invoice["id"],"payment_basis":"FULL_INVOICE"}
    assert signed.put("/api/cases/"+case_id+"/context",json=payload).status_code==200
    process_all()
    detail=signed.get("/api/cases/"+case_id).json()
    report=detail["reports"][0]
    assert report["snapshot"]["sop_packages"]["forex"]["coverage"]["complete"]
    branch=detail["branch_id"]
    response=signed.get("/api/reports/kertas-kerja?branch_id="+branch)
    book=load_workbook(io.BytesIO(response.content))
    assert "Forex SOP Package" in book
    sheet=book["Forex SOP Package"]
    assert sheet.max_row==16 and sheet.cell(2,2).value=="SOP3:1"


def test_invoice_linkage_also_checks_lld_payment_value(sample):
    sample[1][3]["normalized"]["transaction_value"]["amount"]="999.99"
    assert results(sample)["SYS_14_LINK_POLICY"]["result"]=="FAIL"
    sample[1][3]["normalized"]["transaction_value"].update(amount="1000.00",currency="EUR")
    assert results(sample)["SYS_14_LINK_POLICY"]["result"]=="FAIL"


def test_field_number_locale_takes_precedence(sample):
    cfg,docs,ctx=sample
    cfg["profiles"]["application"]["policy"].update(number_locale="id_ID",field_number_locales={"amount":"en_US"})
    docs[1]["normalized"]["amount"]="1,000"
    assert results(sample)["SYS_08_AMOUNT"]["result"]=="PASS"
