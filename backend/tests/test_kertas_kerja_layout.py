import copy
import io
from datetime import date
from types import SimpleNamespace
from openpyxl import load_workbook
from app.domain import default_config
from app.kertas_kerja import workbook
from app.branch_report import select_branch_reports


def sample():
    cfg=default_config()
    scores={"DATA_COMPLETION":0,"CUSTOMER_SIGNATURE":None,"DOCUMENT_COMPLETION":1,"OFFICER_REVIEW_APPROVAL":1,"total_a":0,"total_b":1,"total":0}
    return {"customer":"Test Customer","reference":"CASE-1","configuration":cfg,"template":"forex","provider":"fixture","config_version":1,"generated_at":1700000000,"organization":"sqe","branch_code":"001","branch_name":"Test Branch","validation_status":"FAIL","scores":scores,"findings":[{"rule_id":"INTERNAL_ID","field":"application.notes","source_ref":"SOP3:1","scope_basis":"SYSTEM_SOP3","category":"DATA_COMPLETION","eligible":True,"result":"FAIL","message":"The customer name differs from the identity document.","evidence":{}},{"rule_id":"SIGNATURE","field":"application.customer_signature","category":"CUSTOMER_SIGNATURE","eligible":True,"result":"REVIEW","message":"Inspect both signature images.","evidence":{}}],"documents":[{"id":"app","doc_type":"application","status":"COMPLETED","filename":"test.pdf","sha256":"source","normalized":{"remitter":{"name":"Test Customer"},"remitter_account_number":"0000012345","transaction_date":"10-07-2025"},"raw":{},"overrides":{},"validation":{}}],"case_context":{"transaction_carrier":"CUSTOMER"},"report_export_context":{"period":"01 Jul 2025 – 30 Sep 2025 (transaction dates)"}}


def test_reference_layout_preserves_source_scores_and_values():
    snap=sample();original=copy.deepcopy(snap)
    book=load_workbook(io.BytesIO(workbook([snap])))
    sheet=book["Kertas Kerja"]
    assert {"L10:N10","O10:Q10","A10:A11","T10:T11"} <= {str(r) for r in sheet.merged_cells.ranges}
    assert sheet["T10"].value=="Remarks (Product)"
    assert sheet["L10"].fill.fgColor.rgb=="00666666"
    assert sheet["E12"].value=="0000012345" and sheet["E12"].data_type=="s"
    assert sheet["I12"].value=="Customer Themselves"
    assert sheet["M12"].value=="Pending" and sheet["N12"].value==0
    assert sheet["P12"].value==1 and sheet["R12"].value==0
    assert "=== Data Completion ===" in sheet["T12"].value
    assert "INTERNAL_ID" not in sheet["T12"].value
    assert sheet["T12"].alignment.vertical=="top"
    assert sheet["G12"].fill.fgColor.rgb=="00BFBFBF"
    assert sheet["A12"].border.left.style=="thin"
    assert sheet["D6"].value.startswith("01 Jul 2025")
    assert sheet["L15"].value==0 and sheet["M15"].value=="Not assessed"
    assert sheet["M16"].value==0 and sheet["M17"].value==1
    assert snap==original


def test_long_remarks_have_explicit_detail_reference_and_full_findings():
    snap=sample();snap["findings"]=[]
    for n in range(15):
        snap["findings"].append({"rule_id":f"RULE_{n}","field":"application.notes","eligible":True,"category":"DATA_COMPLETION","result":"FAIL","message":("Long finding description. "*20)+str(n)})
    book=load_workbook(io.BytesIO(workbook([snap])))
    sheet=book["Kertas Kerja"]
    assert "+ 13 more findings" in sheet["T12"].value
    assert "see SOP Details" in sheet["T12"].value
    assert book["SOP Details"].max_row==16
    assert book["SOP Details"]["B16"].value=="RULE_14"
    assert 110 <= sheet.row_dimensions[12].height <=409


def test_account_opening_layout_is_not_given_forex_merged_headers():
    snap=sample();snap["template"]="account_opening"
    snap["configuration"]["templates"]["account_opening"]={"report_layout":"account_opening","label":"Account Opening"}
    book=load_workbook(io.BytesIO(workbook([snap])))
    sheet=book["Account Opening Kertas Kerja"]
    assert "L10:N10" not in {str(r) for r in sheet.merged_cells.ranges}
    assert sheet["U11"].value=="Remarks (Product)"


def test_latest_customer_selection_filters_current_case_version_before_period():
    cases={"one":SimpleNamespace(customer="Customer",template="forex",status="READY"),"two":SimpleNamespace(customer=" customer ",template="forex",status="READY"),"ao":SimpleNamespace(customer="Customer",template="account_opening",status="READY")}
    db=SimpleNamespace(get=lambda _,id:cases[id])
    def report(case,version,created,day):
        snap=sample();snap["documents"][0]["normalized"]["transaction_date"]=day
        return SimpleNamespace(case_id=case,version=version,created_at=created,id=f"{case}-{version}",snapshot=snap)
    reports=[report("one",1,1,"10-07-2025"),report("one",2,2,"10-10-2025"),report("two",1,3,"11-07-2025"),report("ao",1,4,"11-07-2025")]
    current=select_branch_reports(db,reports)
    assert {r.id for r in current.values()}=={"two-1","ao-1"}
    july=select_branch_reports(db,reports,"forex",date(2025,7,1),date(2025,7,31))
    assert {r.id for r in july.values()}=={"two-1"}
    assert not select_branch_reports(db,reports[:2],"forex",date(2025,7,1),date(2025,7,31)) # no fallback to old version


def test_export_api_validates_period_and_deduplicates_customer(signed,monkeypatch):
    from test_latest_reports import seed_report
    from app.db import Case,Report,SessionLocal
    case_id,old,new,branch,config_id,user_id=seed_report(signed)
    with SessionLocal.begin() as db:
        original=db.get(Report,new)
        duplicate=Case(org="sqe",branch_id=branch,customer=" sri p ",reference="duplicate",template="forex",created_by=user_id,config_id=config_id,status="NEEDS_REVIEW")
        db.add(duplicate);db.flush()
        db.add(Report(org="sqe",branch_id=branch,case_id=duplicate.id,version=1,fingerprint="duplicate",snapshot=copy.deepcopy(original.snapshot),created_at=original.created_at+10))
    captured=[]
    monkeypatch.setattr("app.kertas_kerja.workbook",lambda snapshots:captured.extend(snapshots) or b"fixture-workbook")
    response=signed.get("/api/reports/kertas-kerja",params={"branch_id":branch})
    assert response.status_code==200 and len(captured)==1
    assert captured[0]["report_export_context"]["period"]=="All transaction dates (unfiltered)"
    assert signed.get("/api/reports/kertas-kerja",params={"branch_id":branch,"date_from":"2025-09-01","date_to":"2025-07-01"}).status_code==422
    assert signed.get("/api/reports/kertas-kerja",params={"branch_id":branch,"date_from":"invalid"}).status_code==422
    assert signed.get("/api/reports/kertas-kerja",params={"branch_id":branch,"transaction_flow":"not-present"}).status_code==404
    assert signed.get("/api/reports/kertas-kerja",params={"branch_id":branch,"date_from":"2025-01-01"}).status_code==404 # no readable dates


def test_historical_data_review_is_not_labelled_visual_review():
    snap=sample()
    snap["findings"][0]["result"]="REVIEW"
    snap["findings"][1]["result"]="PASS"
    snap["scores"].update(DATA_COMPLETION=None,CUSTOMER_SIGNATURE=1,total_a=None,total=None)
    sheet=load_workbook(io.BytesIO(workbook([snap])))["Kertas Kerja"]
    assert sheet["L12"].value=="Not assessed"
    assert "Awaiting data assessment" in sheet["T12"].value
    assert "Awaiting visual review: 0" in sheet["T12"].value


def test_pending_indicator_is_scoped_to_its_category_group():
    from app.forex_report_layout import display_score
    snap=sample();snap["scores"].update(total_a=None,total_b=None,total=None)
    assert display_score(snap,"total_a")=="Pending"
    assert display_score(snap,"total_b")=="Not assessed"
    assert display_score(snap,"total")=="Pending"
