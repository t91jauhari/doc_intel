"""Export final results from durable field requests, including failed/unscored rows."""
import argparse,csv,json,sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
import openpyxl
from sqlalchemy import select
from app.db import SessionLocal,FieldReprocess,Document,Case,Job
from app.reporting import csv_safe
from validation_inventory import inventory
from run_field_regression import norm,text,score

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('workbook');p.add_argument('--output',required=True);a=p.parse_args()
    root=Path(a.output);root.mkdir(parents=True,exist_ok=True)
    book=openpyxl.load_workbook(a.workbook,data_only=True);targets=inventory(a.workbook);rows=[]
    with SessionLocal() as db:
        requests=list(db.scalars(select(FieldReprocess).order_by(FieldReprocess.created_at.desc())))
        requests=[r for r in requests if r.payload.get('reason','').startswith('Validation workbook')]
        jobs=[]
        for req in requests:
            job=db.get(Job,req.job_id);doc=db.get(Document,req.document_id)
            jobs.append({'request_id':req.id,'document_id':doc.id,'filename':doc.filename,'status':job.status,'error':job.error,
                         'fields':req.payload['fields'],'config_version':req.payload['config_version']})
        for target in targets:
            match=None
            for req in requests:
                doc=db.get(Document,req.document_id);case=db.get(Case,doc.case_id)
                if doc.doc_type!=target['profile'] or target['field'] not in req.payload['fields']:continue
                if norm(case.customer).startswith(norm(target['sheet'])) or norm(target['sheet']).startswith(norm(case.customer)):
                    match=(req,doc,db.get(Job,req.job_id));break
            source=book[target['sheet']].cell(target['row'],3).value
            previous=book[target['sheet']].cell(target['row'],4).value
            if match:
                req,doc,job=match;actual=doc.normalized.get(target['field'])
                result=score(source,actual,target['field']) if job.status=='COMPLETED' else 'Unable to verify'
                note='Re-extracted from original document; compared with validation workbook reference.' if job.status=='COMPLETED' else job.error
                request_id=req.id;doc_id=doc.id
            else:actual=None;result='Unable to verify';note='No matching field request';request_id='';doc_id=''
            if target['sheet']=='Sri P' and target['profile']=='riplay' and target['field'] in ['bank_officer_name','customer_name']:
                note+=' Reference discrepancy confirmed visually: the original signing block contains the handwritten names; source workbook says blank.'
            if target['sheet']=='Indah' and target['field']=='remitter_account_number':
                note+=' Reference formatting issue: the validation workbook lost a leading zero; Forex-Dataset-Sample.xlsx Sample 1 retains 0800997624.'
            if target['sheet']=='Steve' and target['field']=='shipping_handling':
                note+=' Reference discrepancy: CIF JAKARTA is visibly printed on the original invoice; validation source is N/A.'
            rows.append({'Document Name':target['document'],'Customer':target['sheet'],'Field Name':target['label'],
                'Source Value':text(source),'Previous System Value':text(previous),'System Value':text(actual),
                'Previous Comparison Result':target['status'],'Comparison Result':result,'Notes':note,
                'Profile Field Key':target['field'],'Document ID':doc_id,'Request ID':request_id})
    summary=[]
    for customer,document in sorted({(r['Customer'],r['Document Name']) for r in rows})+[('Overall','All targeted invalid rows')]:
        group=rows if customer=='Overall' else [r for r in rows if r['Customer']==customer and r['Document Name']==document]
        valid=sum(r['Comparison Result']=='Valid (same result)' for r in group)
        invalid=sum(r['Comparison Result']=='Invalid (different value)' for r in group)
        unknown=len(group)-valid-invalid
        summary.append({'Document Name':document,'Customer':customer,'Previously Invalid Rows':len(group),
            'Reference Matches After':valid,'Remaining Differences':invalid,'Unable to Verify':unknown,
            'Reference Match Rate on Reprocessed Rows':f'{100*valid/(valid+invalid):.2f}%' if valid+invalid else 'N/A',
            'Scope':'Previously invalid rows only; not whole-dataset accuracy'})
    for filename,data in [('Extraction_Field_Regression.csv',rows),('Extraction_Field_Regression_Summary.csv',summary)]:
        with (root/filename).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=list(data[0]),quoting=csv.QUOTE_ALL);writer.writeheader()
            writer.writerows({k:csv_safe(v) for k,v in row.items()} for row in data)
    (root/'field-jobs.json').write_text(json.dumps(jobs,indent=2),encoding='utf-8')
    (root/'regression-summary.json').write_text(json.dumps(summary[-1],indent=2),encoding='utf-8')
    print(json.dumps({'requests':len(jobs),'completed':sum(j['status']=='COMPLETED' for j in jobs),'errors':sum(j['status']!='COMPLETED' for j in jobs),**summary[-1]}))
